import io,zipfile
import pytest
from docx import Document
from pypdf import PdfWriter
from app.parsing import parse_file
@pytest.mark.parametrize('data,name',[(b'', 'empty.pdf'),(b'not pdf','x.pdf'),(b'hello','x.exe'),(b'PKwrong','x.docx')])
def test_invalid_upload(data,name):
    with pytest.raises(ValueError):parse_file(data,name)
def test_oversize():
    with pytest.raises(ValueError,match='5 MB'):parse_file(b'x'*(5*1024*1024+1),'x.pdf')
def test_docx_paragraph_and_table():
    doc=Document();doc.add_paragraph('Skills');doc.add_paragraph('Python and SQL');t=doc.add_table(rows=1,cols=2);t.cell(0,0).text='Projects';t.cell(0,1).text='Built predictive models';b=io.BytesIO();doc.save(b)
    text=parse_file(b.getvalue(),'resume.docx');assert 'Python' in text and 'Built predictive models' in text

def test_empty_pdf():
    writer=PdfWriter();writer.add_blank_page(width=300,height=300);b=io.BytesIO();writer.write(b)
    with pytest.raises(ValueError,match='No selectable text'):parse_file(b.getvalue(),'scan.pdf')
def test_archive_expansion_limit():
    b=io.BytesIO()
    with zipfile.ZipFile(b,'w',zipfile.ZIP_DEFLATED) as z:z.writestr('word/document.xml',' '*(21*1024*1024))
    with pytest.raises(ValueError,match='expands'):parse_file(b.getvalue(),'bomb.docx')
def test_txt_job_only():
    assert 'Python' in parse_file(b'Required Python and SQL','job.txt',True)
    with pytest.raises(ValueError):parse_file(b'Required Python','resume.txt')
