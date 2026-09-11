import io
from docx import Document
from app.main import limits

def test_health_and_auth(client):
    assert client.get('/health').status_code==200
    assert client.get('/recommendations?resume_id=unknown').status_code==401
    assert client.get('/skills').status_code==200

def test_full_workflow(client,auth):
    r=client.post('/demo/profile',headers=auth);assert r.status_code==201;rid=r.json()['id']
    j=client.post('/job/analyze',headers=auth,json={'text':'Data Scientist\nRequired: Python, SQL, AWS\nResponsibilities: Build predictive models.'});assert j.status_code==201
    result=client.post('/match',headers=auth,json={'resume_id':rid,'job_id':j.json()['id']});assert result.status_code==200;assert 'AWS' in result.json()['missing_skills']
    rec=client.get('/recommendations',headers=auth,params={'resume_id':rid});assert rec.status_code==200;assert len(rec.json()['jobs'])==19
    analytics=client.get('/market/analytics',headers=auth);assert analytics.status_code==200;assert analytics.json()['sample_count']==18

def test_session_isolation(client,auth):
    rid=client.post('/demo/profile',headers=auth).json()['id'];j=client.post('/job/analyze',headers=auth,json={'text':'Private job\nRequired: Python and SQL skills.'}).json()
    other={'Authorization':'Bearer '+client.post('/session').json()['token']}
    assert client.get('/resumes',headers=other).json()==[]
    assert client.get('/recommendations',headers=other,params={'resume_id':rid}).status_code==404
    assert j['id'] not in [x['id'] for x in client.get('/jobs',headers=other).json()]
    assert client.delete('/resume/'+rid,headers=other).status_code==404

def test_upload_docx(client,auth):
    d=Document();d.add_paragraph('Skills');d.add_paragraph('Python, SQL, Pandas');b=io.BytesIO();d.save(b)
    r=client.post('/resume/upload',headers=auth,files={'file':('resume.docx',b.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')});assert r.status_code==201;assert 'Python' in r.json()['structured']['skills']

def test_invalid_file_and_empty(client,auth):
    assert client.post('/resume/upload',headers=auth,files={'file':('resume.pdf',b'fake','application/pdf')}).status_code==422
    assert client.post('/resume/text',headers=auth,json={'text':' '*40}).status_code==422
    assert client.post('/job/analyze',headers=auth,json={'text':''}).status_code==422

def test_unusual_job(client,auth):
    r=client.post('/job/analyze',headers=auth,json={'text':'We seek someone curious, thoughtful, and willing to collaborate.'});assert r.status_code==201;assert r.json()['structured']['required_skills']==[]

def test_delete_cascades(client,auth):
    rid=client.post('/demo/profile',headers=auth).json()['id']
    client.post('/match',headers=auth,json={'resume_id':rid,'job_id':'demo-01'})
    assert client.delete('/resume/'+rid,headers=auth).status_code==204
    assert client.get('/resumes',headers=auth).json()==[]
    assert client.delete('/session',headers=auth).status_code==204
    assert client.get('/resumes',headers=auth).status_code==401

def test_rate_limit(client):
    for _ in range(120):r=client.get('/health')
    assert client.get('/health').status_code==429


def test_repeated_match_upsert(client,auth):
    rid=client.post('/demo/profile',headers=auth).json()['id']
    for _ in range(2):
        assert client.post('/match',headers=auth,json={'resume_id':rid,'job_id':'demo-01'}).status_code==200
        assert client.get('/recommendations',headers=auth,params={'resume_id':rid,'role':'Data Scientist'}).status_code==200

def test_long_role_rejected(client,auth):
    assert client.get('/jobs',headers=auth,params={'role':'x'*101}).status_code==422

def test_chunked_oversize(client,auth):
    def body():
        for _ in range(7):yield b'x'*(1024*1024)
    assert client.post('/resume/upload',headers={**auth,'Content-Type':'application/octet-stream'},content=body()).status_code==413
